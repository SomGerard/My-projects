#!Python 3 automateResultLetter.py
#This script reads the results of a class from an excel file and
#creates result letters for pupils of the class into a word document
# with one page representing the result of a single person
#created by Akuekwe Somtochukwu Emmanuel
#project started on Monday jul 6th, 2020
#project completed; Tuesday jul 28th, 2020

import docx,pprint,openpyxl,re,sys,writeDocx
wb = openpyxl.load_workbook(sys.argv[1])       #workbook to read students data from 
sheet=wb.active                                      #sheet to read from
docxFileName = sys.argv[2]                            #docx file that contains the template sample for writing the letter
runTimes= writeDocx.scriptRunTimes               #get the number of times this script has been executed on a particular computer and use it to create documents
textFileName = 'student letters_' + str(runTimes) + '.txt'              #name of text document to read in the edited data
resultFileName = 'student letters_' + str(runTimes) + '.docx'            #name of resulting docx file
        
    
#gets the subjects and grades data for a particular student
#this helps to arrange the results for each student
# as a dictionary  is an unsorted data structure
def getSubjects(index, firstName, lastName, fullName):
    subjects = {'dict' : 0}
    row = index
    if fullName != (firstName +' '+lastName):
        subjects.clear()
    while((sheet['A'+ str(row)].value==firstName) and (sheet['B'+ str(row)].value==lastName)):
        subjectName = sheet['D' + str(row)].value
        grade = sheet['E' + str(row)].value
        subjects.setdefault(subjectName, grade)
        row+=1
    return subjects
    
#organizes the data for each student into a dictionary structure
def structureData():
    studentsData = {}
    fullName = ' '
    for row in range(2, sheet.max_row+1):
        firstName = sheet['A' + str(row)].value
        lastName =sheet['B' +str(row)].value
        sex = sheet['C' +str(row)].value
        courses = getSubjects(row, firstName, lastName, fullName)
        fullName = firstName + ' ' + lastName
        studentsData.setdefault(firstName, {'lastname': lastName, 'gender': sex, 'subjects': courses})
    return studentsData

#this function creates a list of students first names
#to ensure that the data is arranged sequentially
def createNamesList():
    name_list = []
    firstName = sheet['A2'].value
    lastName = sheet['B2'].value
    fullName = firstName+' '+lastName
    name_list.append(firstName)
    limit = sheet.max_row+1
    for row in range(2, limit):
        firstName = str(sheet['A'+str(row)].value)
        lastName = str(sheet['B' +str(row)].value)
        if fullName == firstName+' '+lastName:
            continue
        else:
            if firstName== 'None':
                break
            else:
                name_list.append(firstName)
        fullName= firstName+' '+lastName
    return name_list
    

#read the full stream of text from the sample letter into a  list
def getText(filename):
    doc = docx.Document(filename)
    fullText =[]
    for para in doc.paragraphs:
        fullText.append(str(para.text))
    for line in fullText:
        if line == '':         #remove extra paragraph space read from the docx file
            fullText.remove(line)
    return fullText

#writeText take a list of strings and write into a text file
def writeText(textList, text_file):
    for line in textList:
        if line == ' ':
            textList.remove(line)
    for line in textList:
        text_file.write(str(line))
        text_file.write('\n')

#choose comment to include in the letter for a particular student depending on their grade
#this function takes two arguments, the dictionary that maps
#the subjects taken by a particular student to their grade in the subject, and
#the list of text to be written in the letter for the particular student
def writeComment(subjects_dict, text_list):
    red = 0
    amber = 0
    green = 0
    none =0
    for key in subjects_dict.keys():
        if subjects_dict[key] == 'Red':
            red+=1
        elif subjects_dict[key] == 'Amber':
            amber+=1
        elif  subjects_dict[key] == 'Green':
            green+=1
        else:
            none+=1
    if red >= amber+green:
        text_list.pop(16)
    else:
        text_list.pop(14)
        text_list.pop(14)
        
        
        
        
    
#udateText; look through the text returned by getText
#and customize letters for each student
def updateText(name_list, data, letter_text):
    textFile = open(textFileName, 'a') 
    for i in range(7, 13):
        letter_text[i] = ' '
    letter_text.insert(13, ' ')
    text_buffer =letter_text.copy()
    for i in range (0, len(name_list)):
        for j in range (0, len(letter_text)):
            letter_text[j] = re.sub(r'<FirstName>', name_list[i], letter_text[j])
            letter_text[j] = re.sub(r'<Lastname>', str(data[name_list[i]]['lastname']), letter_text[j])
            if data[name_list[i]]['gender'] == 'Male':
                letter_text[j] = re.sub(r'<his/her>','his', letter_text[j])
                letter_text[j] = re.sub(r'<him/her>', 'him', letter_text[j])
            else:
                letter_text[j] = re.sub(r'<his/her>','her', letter_text[j])
                letter_text[j] =re.sub(r'<him/her>', 'her', letter_text[j])
        index=7
        for key in data[name_list[i]]['subjects'].keys():
            if data[name_list[i]]['subjects'][key] == None:
                continue
            else:
                letter_text[index] = str(key)+' = '+data[name_list[i]]['subjects'][key]
                index+=1
        writeComment(data[name_list[i]]['subjects'], letter_text)
        letter_text.append('end\n\n')
        writeText(letter_text, textFile)
        letter_text = text_buffer.copy()
    textFile.close()
#write the sudents letters from the text file into a word document   


#this is the main program
def main():
    data=structureData()                                  #organize the data from the excel sheet into a dictionary data structure
    namesList = createNamesList()                         #create a list of first names for sequential data accessing
    letterText = getText(docxFileName)                    #read text from word document template and organize into a list of strings
    print('Text sussefully copied!!!!\n\n Writing text document.......')             
    #print(letterText)
    updateText(namesList, data, letterText)               #update the text in the list and write into a .txt file
    print('Text file generated!!\n')
    print('Getting input from text file...........\n\n')
    inputText = writeDocx.seperateStrings(textFileName)
    print('Writing students letters...............\n\n')
    writeDocx.writeDocx(inputText, resultFileName)
    print('Completed!')
    
    
        
#top level codeexecutes here
if __name__=="__main__":
    main()
    #update number of times the program has been run
    runTimes+=1
    fileObj = open('writeDocx.py', 'a')
    fileObj.write('scriptRunTimes='+pprint.pformat(runTimes)+ '\n')
    fileObj.close()
    

